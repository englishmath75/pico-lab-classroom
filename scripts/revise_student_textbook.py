from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT.parent / "upload" / "arduino-uno-student-textbook.docx"
OUTPUT = ROOT / "public" / "downloads" / "arduino-uno-student-textbook.docx"
RESISTOR_IMAGE = ROOT / "public" / "downloads" / "resistor-reading-guide.png"
NAVY, CYAN, AMBER, PALE, LINE = "0F172A", "0891B2", "FBBF24", "F8FAFC", "D9E2EC"

def font(run, size=11, bold=False, color=NAVY, name="Noto Sans KR"):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr(); rf = rpr.rFonts
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for key, value in (("ascii", name), ("hAnsi", name), ("eastAsia", "Noto Sans KR")):
        rf.set(qn("w:" + key), value)

def para(doc, text="", size=11, bold=False, color=NAVY, align=None, before=0, after=6, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    if align is not None: p.alignment = align
    if text: font(p.add_run(text), size, bold, color)
    return p

def shade(cell, fill):
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); cell._tc.get_or_add_tcPr().append(shd)

def borders(table, color=LINE):
    pr=table._tbl.tblPr; node=OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        x=OxmlElement("w:"+edge); x.set(qn("w:val"),"single"); x.set(qn("w:sz"),"6"); x.set(qn("w:color"),color); node.append(x)
    pr.append(node)

def cell_text(cell, text, size=10.5, bold=False, color=NAVY, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text=""; p=cell.paragraphs[0]; p.alignment=align; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.15
    font(p.add_run(text),size,bold,color); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def heading(doc, text, level=1):
    p=para(doc, text, size=17 if level==1 else 13, bold=True, color="000000", before=12, after=8, style=f"Heading {level}")
    p.paragraph_format.keep_with_next=True; return p

def category(doc, title, definition, inputs, outputs, commands, fill):
    t=doc.add_table(rows=2, cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; borders(t)
    top=t.cell(0,0); top.merge(t.cell(0,2)); shade(top,fill)
    cell_text(top,title+"\n"+definition,15,True,"FFFFFF",WD_ALIGN_PARAGRAPH.CENTER)
    labels=[("입력 부품과 센서",inputs),("출력 부품",outputs),("핵심 명령어",commands)]
    for i,(label,body) in enumerate(labels):
        shade(t.cell(1,i),"FFFFFF")
        p=t.cell(1,i).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
        font(p.add_run(label+"\n"),10,True,fill); font(p.add_run(body),11.5,True,NAVY,"Consolas" if i==2 else "Arial")
    para(doc,"",after=2)

def add_new_front(doc):
    para(doc,"Arduino Uno",25,True,"000000",WD_ALIGN_PARAGRAPH.CENTER,after=3,style="Title")
    para(doc,"디지털 아날로그 시리얼로 시작하는 학생용 실습 교재",15,True,CYAN,WD_ALIGN_PARAGRAPH.CENTER,after=8)
    para(doc,"이름 ____________________   학번 __________   모둠 __________",10.5,False,"475569",WD_ALIGN_PARAGRAPH.CENTER,after=14)
    para(doc,"아두이노 코드는 세 가지 통로를 구분하면 쉽게 읽을 수 있습니다. 무엇을 입력받고, 어떤 명령으로 처리하며, 어떤 부품으로 출력하는지 먼저 찾으세요.",11.5,False,NAVY,after=10)
    category(doc,"1  디지털","0 또는 1, LOW 또는 HIGH처럼 두 상태를 다룹니다.","버튼\n스위치\n디지털 센서","LED\n부저\n릴레이","pinMode()\ndigitalRead()\ndigitalWrite()",NAVY)
    category(doc,"2  아날로그","여러 단계로 변하는 값을 읽고 PWM 출력으로 바꿉니다.","가변저항\nCDS 조도센서\n아날로그 센서","PWM LED 밝기\n모터 속도\n부저 음높이","analogRead()\nmap()\nanalogWrite()","B45309")
    category(doc,"3  시리얼","USB로 컴퓨터와 Arduino가 문자와 숫자를 주고받습니다.","키보드 입력\n시리얼 모니터\n다른 장치의 문자","센서값 표시\n상태 메시지\n디버깅 정보","Serial.begin()\nSerial.available()\nSerial.read()\nSerial.print()",CYAN)
    para(doc,"핵심 흐름  입력 → 처리 → 출력",14,True,"000000",WD_ALIGN_PARAGRAPH.CENTER,before=8,after=4)
    para(doc,"배선을 바꿀 때는 USB를 분리하고, LED에는 220Ω 또는 330Ω 저항을 직렬로 연결합니다.",9.5,True,"B91C1C",WD_ALIGN_PARAGRAPH.CENTER,after=0)
    doc.add_page_break()

COLORS={"검정":"111827","갈색":"7C2D12","빨강":"DC2626","주황":"F97316","노랑":"FACC15","초록":"16A34A","파랑":"2563EB","보라":"7E22CE","회색":"6B7280","흰색":"F8FAFC","금색":"D4A017","은색":"B8C2CC"}
def make_resistor_image():
    W,H=1700,620; im=Image.new("RGB",(W,H),"white"); d=ImageDraw.Draw(im)
    try: f=ImageFont.truetype(str(ROOT / "runtime-fonts" / "NotoSansKR-Regular.ttf"),27); fb=ImageFont.truetype(str(ROOT / "runtime-fonts" / "NotoSansKR-Bold.ttf"),38)
    except: f=fb=None
    rows=[("220Ω",["빨강","빨강","갈색","금색"]),("330Ω",["주황","주황","갈색","금색"]),("1kΩ",["갈색","검정","빨강","금색"]),("10kΩ",["갈색","검정","주황","금색"])]
    for i,(label,bands) in enumerate(rows):
        y=45+i*140; d.text((45,y+30),label,fill="#"+NAVY,font=fb)
        d.line((230,y+55,360,y+55),fill="#475569",width=8); d.line((1040,y+55,1140,y+55),fill="#475569",width=8)
        d.rounded_rectangle((350,y,1050,y+110),radius=42,fill="#F4E7C5",outline="#334155",width=5)
        xs=[440,565,690,900]
        for x,c in zip(xs,bands): d.rectangle((x,y+5,x+48,y+105),fill="#"+COLORS[c],outline="#111827",width=2)
        d.text((1170,y+34)," · ".join(bands),fill="#334155",font=f)
    im.save(RESISTOR_IMAGE)

def add_resistor_section(doc):
    doc.add_page_break(); heading(doc,"저항 색띠 읽는 법",1)
    para(doc,"4색 띠 저항은 금색 또는 은색 오차 띠를 오른쪽에 두고 왼쪽부터 읽습니다. 첫째 띠와 둘째 띠는 숫자, 셋째 띠는 곱하는 수, 넷째 띠는 오차 범위입니다.",11.5)
    t=doc.add_table(rows=13,cols=5); t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t)
    for j,x in enumerate(["색","첫째 숫자","둘째 숫자","곱하는 수","오차"]): shade(t.cell(0,j),NAVY); cell_text(t.cell(0,j),x,10,True,"FFFFFF",WD_ALIGN_PARAGRAPH.CENTER)
    rows=[("검정","0","0","×1","-"),("갈색","1","1","×10","±1%"),("빨강","2","2","×100","±2%"),("주황","3","3","×1,000","-"),("노랑","4","4","×10,000","-"),("초록","5","5","×100,000","±0.5%"),("파랑","6","6","×1,000,000","±0.25%"),("보라","7","7","×10,000,000","±0.1%"),("회색","8","8","×100,000,000","±0.05%"),("흰색","9","9","×1,000,000,000","-"),("금색","-","-","×0.1","±5%"),("은색","-","-","×0.01","±10%")]
    for i,row in enumerate(rows,1):
        for j,x in enumerate(row):
            if i%2==0: shade(t.cell(i,j),PALE)
            cell_text(t.cell(i,j),x,9.4,j==0,NAVY,WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break(); heading(doc,"자주 사용하는 저항 네 가지",1)
    make_resistor_image(); doc.add_picture(str(RESISTOR_IMAGE),width=Inches(7.0)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    para(doc,"읽기 예시  빨강 빨강 갈색은 22×10=220Ω입니다. 갈색 검정 주황은 10×1,000=10,000Ω, 즉 10kΩ입니다. 금색 띠는 보통 ±5% 오차를 뜻합니다.",11.5,False,NAVY,before=8)

DESC=[
"Arduino 프로그램에서 setup()과 loop()가 각각 언제 실행되는지 비교하여 설명하시오.",
"디지털 신호와 아날로그 신호의 차이를 버튼과 가변저항을 예로 들어 설명하시오.",
"LED에 220Ω 또는 330Ω 저항을 직렬로 연결해야 하는 이유를 전류의 관점에서 설명하시오.",
"INPUT_PULLUP을 사용할 때 버튼을 누른 상태가 LOW가 되는 이유와 이를 조건문에서 처리하는 방법을 설명하시오.",
"시리얼 통신에서 Serial.available()을 먼저 확인한 뒤 Serial.read()를 실행하는 이유를 설명하시오.",
"analogRead()의 0~1023 값을 analogWrite()의 0~255 값으로 바꾸어야 하는 이유를 설명하시오.",
"PWM이 실제 아날로그 전압을 계속 출력하는 방식이 아닌데도 LED 밝기가 중간 밝기로 보이는 원리를 설명하시오.",
"초음파센서 거리 계산식에서 왕복 시간을 2로 나누는 이유를 설명하시오.",
"CDS 센서의 기준값을 친구의 숫자 그대로 사용하지 않고 자기 회로에서 측정해 정해야 하는 이유를 설명하시오.",
"여러 센서를 통합할 때 각각을 단독으로 먼저 시험해야 하는 이유와 오류를 찾는 순서를 설명하시오."]
SHORT=[
"Arduino가 시작할 때 한 번만 실행하는 함수는 (              )이다.",
"전원이 켜진 동안 계속 반복되는 함수는 (              )이다.",
"디지털 출력 핀을 준비하는 문장은 pinMode(핀, (              ));이다.",
"핀을 HIGH 또는 LOW로 출력하는 함수는 (              )이다.",
"디지털 핀의 상태를 읽는 함수는 (              )이다.",
"Uno의 아날로그 입력 범위는 (       )부터 (       )까지이다.",
"Uno의 PWM 출력값 범위는 (       )부터 (       )까지이다.",
"시리얼 통신 속도를 9600으로 시작하는 문장은 (                         )이다.",
"문자 한 개를 시리얼에서 읽는 함수는 (                         )이다.",
"빨강 빨강 갈색 금색 저항의 저항값은 (              )Ω이다."]
CODE=[
"D3을 출력으로 준비하도록 빈칸을 완성하시오.   pinMode(3, __________);",
"D3 LED를 켜도록 함수 전체를 작성하시오.   ______________________________;",
"D3 LED를 끄도록 함수 전체를 작성하시오.   ______________________________;",
"버튼이 연결된 D2의 값을 변수 a에 저장하시오.   int a = ______________________;",
"A4의 아날로그 값을 변수 a에 저장하시오.   int a = ______________________;",
"a의 0~1023 범위를 0~255로 바꾸어 b에 저장하시오.   int b = ______________________________;",
"D3 LED에 PWM 값 b를 출력하시오.   ______________________________;",
"시리얼 데이터가 있을 때만 읽도록 조건을 완성하시오.   if (________________________ > 0) { char c = Serial.read(); }",
"버튼값 a가 LOW이면 LED를 켜도록 조건과 출력문을 완성하시오.   if (__________) { ____________________________; }",
"1부터 5까지 출력하도록 for문을 완성하시오.   for (int a=1; __________; __________) { Serial.println(a); }"]

def add_question_group(doc,title,intro,questions,lines=2):
    doc.add_page_break(); heading(doc,title,1); para(doc,intro,10.5,False,"475569",after=10)
    for i,q in enumerate(questions,1):
        p=para(doc,f"{i}. {q}",10.8,True,NAVY,before=5,after=4); p.paragraph_format.keep_with_next=True
        for _ in range(lines): para(doc,"답  ____________________________________________________________________________",9.5,False,"64748B",after=4)

def add_assessment(doc):
    add_question_group(doc,"서술형 문제 10문제","이론과 작동 원리를 완전한 문장으로 설명하세요. 핵심 용어만 나열하지 않습니다.",DESC,2)
    add_question_group(doc,"단답형 문제 10문제","괄호 또는 빈칸에 정확한 용어, 숫자, 문장을 쓰세요.",SHORT,1)
    add_question_group(doc,"코드 문제 10문제","앞뒤 코드의 논리를 읽고 들어갈 코드를 직접 작성하세요. 함수 이름의 대소문자와 괄호, 세미콜론까지 확인하세요.",CODE,2)

def main():
    doc=Document(SOURCE); body=doc._element.body
    marker=None
    for p in doc.paragraphs:
        if p.text.strip()=="이 교재로 공부하는 방법": marker=p._element; break
    if marker is None: raise RuntimeError("도입부 기준 제목을 찾지 못했습니다.")
    for el in list(body):
        if el is marker: break
        if el.tag != qn("w:sectPr"): body.remove(el)
    front=Document(); add_new_front(front)
    for el in list(front._element.body):
        if el.tag != qn("w:sectPr"): marker.addprevious(deepcopy(el))
    add_resistor_section(doc); add_assessment(doc)
    doc.core_properties.title="Arduino Uno 학생용 실습 교재"
    doc.core_properties.subject="디지털 아날로그 시리얼 저항 실습 평가"
    doc.core_properties.author="박영수 정보T"
    doc.save(OUTPUT); print(OUTPUT)

if __name__=="__main__": main()
